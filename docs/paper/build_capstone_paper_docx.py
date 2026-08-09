from __future__ import annotations

import re
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
SOURCE_MD = ROOT / "danzzan-high-concurrency-ticketing-capstone-paper-v3.md"
OUTPUT_DOCX = ROOT / "danzzan-high-concurrency-ticketing-capstone-paper-v3.docx"
DOC_TITLE = "DANZZAN 고동시성 티켓팅 시스템 캡스톤 결과 논문"
DOC_SHORT_TITLE = "DANZZAN 캡스톤 결과 논문"

CONTENT_WIDTH_DXA = 9360
CONTENT_WIDTH_IN = 6.5
BODY_FONT = "맑은 고딕"
LATIN_FONT = "Calibri"
CODE_FONT = "Consolas"
HEADING_BLUE = RGBColor(0x2E, 0x74, 0xB5)
HEADING_DARK = RGBColor(0x1F, 0x4D, 0x78)
MUTED = RGBColor(0x66, 0x66, 0x66)
TABLE_HEADER_FILL = "F4F6F9"
TABLE_BORDER = "B8C4D2"


@dataclass(frozen=True)
class DiagramSpec:
    kind: str
    title: str


DIAGRAMS = [
    DiagramSpec("architecture", "DANZZAN 전체 시스템 아키텍처"),
    DiagramSpec("dualgate", "제안하는 이중 게이트 티켓팅 구조"),
    DiagramSpec("state", "사용자 상태 전이와 ACTIVE lease"),
    DiagramSpec("erd", "티켓팅 중심 ERD"),
    DiagramSpec("queue", "대기열 게이트의 Lua 원자 처리"),
    DiagramSpec("async", "Redis claim 이후 최종 영속화 구조"),
    DiagramSpec("sequence", "티켓팅 핵심 시퀀스"),
]


def set_run_font(run, name=BODY_FONT, size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT if name == BODY_FONT else name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT if name == BODY_FONT else name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=8, line=1.333):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_field(paragraph, instruction: str, fallback: str = ""):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction

    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")

    fallback_run = OxmlElement("w:r")
    fallback_text = OxmlElement("w:t")
    fallback_text.text = fallback
    fallback_run.append(fallback_text)

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_separate)
    run._r.append(fallback_run)
    run._r.append(fld_end)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(paragraph, "PAGE", "1")


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, col_widths: list[int]):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(col_widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    existing_grid = tbl.find(qn("w:tblGrid"))
    if existing_grid is not None:
        tbl.remove(existing_grid)
    grid = OxmlElement("w:tblGrid")
    for width in col_widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(1, grid)

    for row in table.rows:
        for cell, width in zip(row.cells, col_widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), TABLE_BORDER)


def set_update_fields_on_open(doc: Document):
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def configure_document(doc: Document):
    apply_page_setup(doc.sections[0])

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    heading_tokens = [
        ("Heading 1", 16, HEADING_BLUE, 18, 10),
        ("Heading 2", 13, HEADING_BLUE, 12, 6),
        ("Heading 3", 12, HEADING_DARK, 8, 4),
    ]
    for style_name, size, color, before, after in heading_tokens:
        style = styles[style_name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208

    set_update_fields_on_open(doc)
    doc.core_properties.title = DOC_TITLE
    doc.core_properties.subject = "캡스톤 논문"
    doc.core_properties.author = "DANZZAN Team"


def apply_page_setup(section):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(22)
    section.bottom_margin = Mm(22)
    section.left_margin = Mm(22.5)
    section.right_margin = Mm(22.5)
    section.header_distance = Mm(12.5)
    section.footer_distance = Mm(12.5)


def start_new_page_section(doc: Document):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_page_setup(section)
    return section


def finalize_sections(doc: Document):
    for index, section in enumerate(doc.sections):
        apply_page_setup(section)
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False

        header = section.header.paragraphs[0]
        footer = section.footer.paragraphs[0]
        header.clear()
        footer.clear()

        if index == 0:
            continue

        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run(DOC_SHORT_TITLE)
        set_run_font(run, size=8.5, color=MUTED)
        add_page_number(footer)


def add_cover(doc: Document):
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=70, after=18, line=1.2)
    r = p.add_run("캡스톤 디자인 논문")
    set_run_font(r, size=15, color=MUTED, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=20, after=10, line=1.2)
    r = p.add_run("DANZZAN 고동시성 티켓팅 시스템\n캡스톤 결과 논문")
    set_run_font(r, size=21, color=RGBColor(0x14, 0x2F, 0x4C), bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=6, after=42, line=1.2)
    r = p.add_run("축제 통합 서비스와 Redis Lua 기반 티켓팅 정합성 검증")
    set_run_font(r, size=13, color=HEADING_DARK, bold=True)

    metadata = [
        ("프로젝트명", "DANZZAN - 단국대학교 축제 통합 서비스"),
        ("논문 유형", "캡스톤 디자인 결과 논문"),
        ("소속", "[학과/전공 입력]"),
        ("팀명", "DANZZAN"),
        ("팀원", "[팀원 이름 입력]"),
        ("지도교수", "[지도교수 입력]"),
        ("제출일", "2026년 [월] [일]"),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    set_table_borders(table)
    set_table_geometry(table, [1900, 7460])
    for row, (label, value) in zip(table.rows, metadata):
        label_cell, value_cell = row.cells
        set_cell_shading(label_cell, "EEF3F8")
        for cell, text, bold in ((label_cell, label, True), (value_cell, value, False)):
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_spacing(p, before=0, after=0, line=1.15)
            run = p.add_run(text)
            set_run_font(run, size=10.5, bold=bold, color=RGBColor(0x22, 0x22, 0x22))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=48, after=0, line=1.2)
    r = p.add_run("※ 학교 지정 양식이 있는 경우 표지 정보와 여백만 해당 양식에 맞춰 조정한다.")
    set_run_font(r, size=9, color=MUTED)
    start_new_page_section(doc)


def extract_front_matter(markdown: str) -> tuple[list[str], str, str]:
    abstract_match = re.search(r"## 초록\n\n(.+?)\n\n\*\*주요어:\*\* (.+?)\n\n## 1\. 서론", markdown, re.S)
    if not abstract_match:
        raise ValueError("Abstract section not found")
    abstract_paragraphs = [p.strip() for p in abstract_match.group(1).split("\n\n") if p.strip()]
    keywords = abstract_match.group(2).strip()
    body_start = markdown.index("## 1. 서론")
    body = markdown[body_start:]
    return abstract_paragraphs, keywords, body


def add_abstract(doc: Document, abstract_paragraphs: Iterable[str], keywords: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=16, after=16, line=1.2)
    r = p.add_run("국문 초록")
    set_run_font(r, size=16, color=HEADING_BLUE, bold=True)

    for text in abstract_paragraphs:
        add_rich_paragraph(doc, text)

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=4, after=0, line=1.2)
    r = p.add_run("주요어: ")
    set_run_font(r, size=10.5, bold=True)
    add_inline_runs(p, keywords, size=10.5)
    start_new_page_section(doc)


def add_toc(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=16, after=18, line=1.2)
    r = p.add_run("목차")
    set_run_font(r, size=16, color=HEADING_BLUE, bold=True)

    p = doc.add_paragraph()
    add_field(p, r'TOC \o "1-3" \h \z \u', "목차는 Word에서 문서를 열 때 자동 갱신됩니다.")
    set_paragraph_spacing(p, before=0, after=8, line=1.25)
    start_new_page_section(doc)


def add_inline_runs(paragraph, text: str, size=10.5):
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name=CODE_FONT, size=max(size - 0.5, 8.5), color=RGBColor(0x33, 0x33, 0x33))
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size)


def add_rich_paragraph(doc: Document, text: str, style=None, align=None):
    p = doc.add_paragraph(style=style)
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, before=0, after=8, line=1.333)
    add_inline_runs(p, text)
    return p


def add_heading(doc: Document, text: str, level: int):
    style = f"Heading {min(level, 3)}"
    p = doc.add_paragraph(style=style)
    add_inline_runs(p, text, size={1: 16, 2: 13, 3: 12}.get(level, 12))
    return p


def parse_table_rows(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        if re.match(r"^\|\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?$", line.strip()):
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def table_widths(rows: list[list[str]]) -> list[int]:
    count = len(rows[0])
    if count == 2:
        return [2600, 6760]
    if count == 3:
        return [2100, 2200, 5060]
    if count == 4:
        return [1900, 2500, 1900, 3060]
    if count == 5:
        return [1600, 1700, 1700, 2000, 2360]
    if count == 6:
        return [1500, 1400, 1400, 1500, 1800, 1760]
    return [CONTENT_WIDTH_DXA // count] * count


def add_markdown_table(doc: Document, table_lines: list[str]):
    rows = parse_table_rows(table_lines)
    if not rows:
        return
    widths = table_widths(rows)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    set_table_borders(table)
    set_table_geometry(table, widths)
    for row_idx, cells in enumerate(rows):
        for col_idx, value in enumerate(cells):
            cell = table.cell(row_idx, col_idx)
            cell.text = ""
            if row_idx == 0:
                set_cell_shading(cell, TABLE_HEADER_FILL)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_idx == 0 or len(value) <= 12 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_spacing(p, before=0, after=0, line=1.15)
            add_inline_runs(p, value, size=9 if len(rows[0]) >= 5 else 9.5)
            for run in p.runs:
                if row_idx == 0:
                    run.bold = True
    after = doc.add_paragraph()
    set_paragraph_spacing(after, before=2, after=8, line=1.0)


def draw_centered_text(draw, box, text, font, fill, max_width=None, line_spacing=7):
    x1, y1, x2, y2 = box
    max_width = max_width or (x2 - x1 - 28)
    lines = []
    for raw_line in text.splitlines():
        words = re.split(r"(\s+)", raw_line)
        current = ""
        for part in words:
            candidate = current + part
            if draw.textlength(candidate, font=font) <= max_width or not current:
                current = candidate
            else:
                lines.append(current.strip())
                current = part.strip()
        if current:
            lines.append(current.strip())
    total_h = sum(draw.textbbox((0, 0), line, font=font)[3] for line in lines) + line_spacing * (len(lines) - 1)
    y = y1 + (y2 - y1 - total_h) / 2
    for line in lines:
        w = draw.textlength(line, font=font)
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, font=font, fill=fill)
        y += draw.textbbox((0, 0), line, font=font)[3] + line_spacing


def load_diagram_font(size: int):
    candidates = [
        "/System/Library/Fonts/AppleSDGothicNeo.ttc",
        "/System/Library/Fonts/Supplemental/NotoSansGothic-Regular.ttf",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
    ]
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            try:
                return ImageFont.truetype(str(path), size)
            except OSError:
                continue
    return ImageFont.load_default()


def arrow(draw, start, end, fill="#4B6075", width=4):
    draw.line([start, end], fill=fill, width=width)
    sx, sy = start
    ex, ey = end
    dx, dy = ex - sx, ey - sy
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    size = 16
    points = [
        (ex, ey),
        (ex - ux * size + px * size * 0.45, ey - uy * size + py * size * 0.45),
        (ex - ux * size - px * size * 0.45, ey - uy * size - py * size * 0.45),
    ]
    draw.polygon(points, fill=fill)


def route_arrow(draw, points, fill="#4B6075", width=4):
    if len(points) < 2:
        return
    if len(points) > 2:
        draw.line(points[:-1], fill=fill, width=width)
    arrow(draw, points[-2], points[-1], fill=fill, width=width)


def box(draw, xy, text, font, fill="#F4F6F9", outline="#8AA6C1", text_fill="#1E2E3E", radius=18):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=3)
    draw_centered_text(draw, xy, text, font, text_fill)


def create_canvas(width=1800, height=980):
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    return img, draw


def draw_dualgate(path: Path):
    img, draw = create_canvas(1900, 980)
    font = load_diagram_font(30)
    small = load_diagram_font(25)
    title = load_diagram_font(34)

    draw.text((80, 55), "Admission-before-Claim: 이중 게이트 티켓팅 구조", font=title, fill="#18324A")

    components = [
        (80, 250, 300, 360, "사용자\n동시 진입"),
        (420, 210, 760, 400, "대기열 게이트\nenter_queue.lua\n- 중복 진입 방지\n- seq 발급\n- WAITING 저장"),
        (880, 250, 1180, 360, "Redis Sorted Set\nseq 기반 대기열"),
        (1280, 210, 1620, 400, "Admission Scheduler\n- slot 계산\n- 선두 사용자 승격"),
        (420, 560, 760, 750, "발급 게이트\nACTIVE lease\n- 권한 부여\n- 만료 통제"),
        (880, 560, 1180, 750, "claim_v2.lua\n- ACTIVE 검증\n- 중복 claim 방지\n- stock DECR"),
        (1280, 560, 1620, 750, "최종 영속화\nMySQL / Outbox\nKafka Consumer"),
    ]
    for idx, (x1, y1, x2, y2, text) in enumerate(components):
        fill = "#EAF2F8" if idx in {1, 4, 5} else "#F4F6F9"
        box(draw, (x1, y1, x2, y2), text, small if idx in {1, 4, 5, 6} else font, fill=fill)

    arrow(draw, (300, 305), (420, 305))
    arrow(draw, (760, 305), (880, 305))
    arrow(draw, (1180, 305), (1280, 305))
    arrow(draw, (1450, 400), (590, 560))
    arrow(draw, (760, 655), (880, 655))
    arrow(draw, (1180, 655), (1280, 655))

    box(draw, (1660, 560, 1850, 750), "rollback /\ncompensation", small, fill="#FFF4F4", outline="#CC6C73")
    arrow(draw, (1620, 655), (1660, 655), fill="#9B1C1C")
    draw.text((450, 825), "핵심 불변식: capacity 초과 발급 없음 · 사용자 중복 발급 없음 · ACTIVE 없는 reserve 성공 없음 · 실패 상태 수렴", font=small, fill="#1E2E3E")
    img.save(path)


def draw_architecture(path: Path):
    img, draw = create_canvas()
    font = load_diagram_font(34)
    small = load_diagram_font(28)
    boxes = {
        "user": (80, 110, 360, 220, "사용자\n모바일 브라우저"),
        "fe": (520, 110, 850, 220, "Frontend\nReact + TypeScript"),
        "lb": (1010, 110, 1300, 220, "로드밸런서\nHTTPS 종료"),
        "be": (1430, 110, 1720, 220, "Backend\nSpring Boot"),
        "redis": (1060, 380, 1360, 500, "Redis\nqueue / stock / status"),
        "mysql": (1450, 380, 1720, 500, "MySQL RDS\n티켓 영속 저장"),
        "outbox": (600, 650, 900, 770, "Outbox\n발행 대기"),
        "kafka": (1030, 650, 1330, 770, "Kafka\nissue topic"),
        "consumer": (1450, 650, 1720, 770, "Consumer\n티켓 발급"),
        "prom": (540, 380, 840, 500, "Prometheus\n메트릭 수집"),
        "grafana": (80, 380, 360, 500, "Grafana\n대시보드"),
    }
    for key, (x1, y1, x2, y2, text) in boxes.items():
        box(draw, (x1, y1, x2, y2), text, font if key in {"user", "fe", "be"} else small)
    arrow(draw, (360, 165), (520, 165))
    arrow(draw, (850, 165), (1010, 165))
    arrow(draw, (1300, 165), (1430, 165))
    arrow(draw, (1570, 220), (1210, 380))
    arrow(draw, (1570, 220), (1585, 380))
    route_arrow(draw, [(1430, 205), (930, 205), (930, 710), (900, 710)])
    arrow(draw, (900, 710), (1030, 710))
    arrow(draw, (1330, 710), (1450, 710))
    arrow(draw, (1585, 650), (1585, 500))
    route_arrow(draw, [(1210, 500), (1210, 575), (790, 575), (790, 500)], fill="#6B7F93", width=3)
    route_arrow(draw, [(1585, 500), (1585, 615), (790, 615), (790, 500)], fill="#6B7F93", width=3)
    arrow(draw, (540, 445), (360, 445))
    img.save(path)


def draw_sequence(path: Path):
    img, draw = create_canvas(1800, 1100)
    font = load_diagram_font(28)
    small = load_diagram_font(24)
    lanes = [180, 460, 740, 1020, 1300, 1580]
    labels = ["사용자", "Frontend", "API", "Redis", "MySQL", "Kafka/Outbox"]
    for x, label in zip(lanes, labels):
        box(draw, (x - 95, 60, x + 95, 125), label, small)
        draw.line((x, 125, x, 1010), fill="#D2DAE3", width=3)
    steps = [
        (180, 460, 190, "티켓팅 화면 진입"),
        (460, 740, 270, "GET /tickets/events"),
        (740, 1300, 350, "이벤트/재고 조회"),
        (180, 460, 440, "대기열 입장"),
        (460, 740, 520, "POST /queue/enter"),
        (740, 1020, 600, "enter_queue.lua"),
        (460, 740, 690, "GET /queue/status polling"),
        (740, 1020, 770, "상태 snapshot"),
        (460, 740, 850, "POST /reserve"),
        (740, 1020, 930, "claim_v2.lua"),
        (740, 1580, 1010, "비동기 발급 이벤트"),
    ]
    for x1, x2, y, text in steps:
        arrow(draw, (x1, y), (x2, y))
        draw.text((min(x1, x2) + 18, y - 34), text, font=font, fill="#26384A")
    img.save(path)


def draw_state(path: Path):
    img, draw = create_canvas()
    font = load_diagram_font(30)
    nodes = {
        "NONE": (80, 120, 300, 210),
        "WAITING": (410, 120, 660, 210),
        "ACTIVE": (770, 120, 1020, 210),
        "PROCESSING": (1130, 120, 1430, 210),
        "SUCCESS": (1520, 120, 1740, 210),
        "DONE": (1520, 360, 1740, 450),
        "SOLD_OUT": (770, 360, 1020, 450),
        "ALREADY": (410, 360, 660, 450),
        "FAILED": (1130, 360, 1430, 450),
        "EXPIRED": (770, 620, 1020, 710),
        "CANCELLED": (410, 620, 660, 710),
    }
    for name, xy in nodes.items():
        fill = "#EAF2F8" if name in {"WAITING", "ACTIVE", "SUCCESS"} else "#F7F8FA"
        box(draw, xy, name, font, fill=fill)
    for a, b in [
        ("NONE", "WAITING"), ("WAITING", "ACTIVE"), ("ACTIVE", "PROCESSING"),
        ("PROCESSING", "SUCCESS"), ("SUCCESS", "DONE"), ("ACTIVE", "SOLD_OUT"),
        ("ACTIVE", "ALREADY"), ("PROCESSING", "FAILED"), ("ACTIVE", "EXPIRED"),
        ("WAITING", "CANCELLED"),
    ]:
        ax = (nodes[a][2], (nodes[a][1] + nodes[a][3]) // 2)
        bx = (nodes[b][0], (nodes[b][1] + nodes[b][3]) // 2)
        if nodes[a][1] == nodes[b][1]:
            arrow(draw, ax, bx)
        else:
            arrow(draw, ((nodes[a][0] + nodes[a][2]) // 2, nodes[a][3]), ((nodes[b][0] + nodes[b][2]) // 2, nodes[b][1]))
    img.save(path)


def draw_erd(path: Path):
    img, draw = create_canvas(1900, 1150)
    font = load_diagram_font(25)
    title_font = load_diagram_font(28)

    tables = {
        "USERS": (80, 80, 430, 280, ["id PK", "student_id", "role", "deleted"]),
        "FESTIVAL_EVENTS": (80, 430, 500, 690, ["id PK", "title", "ticketing_status", "total_capacity"]),
        "USER_TICKETS": (700, 120, 1160, 390, ["id PK", "user_id FK", "event_id FK", "status", "ticketing_order", "seq"]),
        "TICKET_QUEUE_ENTRIES": (700, 500, 1160, 790, ["id PK", "event_id FK", "user_id FK", "status", "seq", "lease_until"]),
        "TICKET_ISSUE_REQUESTS": (1350, 120, 1820, 420, ["id PK", "request_id UK", "event_id", "user_id", "status", "remaining_after_claim"]),
        "OUTBOX_EVENTS": (1350, 570, 1820, 830, ["id PK", "aggregate_type", "aggregate_id UK", "topic", "status", "retry_count"]),
    }
    for name, (x1, y1, x2, y2, fields) in tables.items():
        draw.rounded_rectangle((x1, y1, x2, y2), radius=14, fill="#FFFFFF", outline="#8AA6C1", width=3)
        draw.rectangle((x1, y1, x2, y1 + 48), fill="#EAF2F8", outline="#8AA6C1", width=0)
        draw.text((x1 + 18, y1 + 10), name, font=title_font, fill="#18324A")
        y = y1 + 65
        for field in fields:
            draw.text((x1 + 22, y), field, font=font, fill="#2B3A48")
            y += 38
    arrow(draw, (430, 180), (700, 220))
    arrow(draw, (500, 550), (700, 290))
    arrow(draw, (500, 560), (700, 635))
    arrow(draw, (430, 210), (700, 610))
    arrow(draw, (1160, 250), (1350, 260))
    arrow(draw, (1585, 420), (1585, 570))
    img.save(path)


def draw_queue(path: Path):
    img, draw = create_canvas(1700, 1100)
    font = load_diagram_font(30)
    steps = [
        (620, 80, 1080, 170, "queue/enter 요청"),
        (620, 240, 1080, 330, "dedupKey SET NX"),
        (620, 400, 1080, 490, "seqKey INCR"),
        (620, 560, 1080, 650, "stockKey GET"),
        (620, 720, 1080, 810, "queue ZADD"),
        (620, 880, 1080, 970, "quser Hash HSET"),
    ]
    for xy in steps:
        box(draw, xy[:4], xy[4], font)
    for i in range(len(steps) - 1):
        arrow(draw, ((steps[i][0] + steps[i][2]) // 2, steps[i][3]), ((steps[i + 1][0] + steps[i + 1][2]) // 2, steps[i + 1][1]))
    box(draw, (1160, 235, 1550, 335), "이미 진입\nWAITING 응답", font, fill="#FFF8E5", outline="#D5A84A")
    arrow(draw, (1080, 285), (1160, 285))
    box(draw, (1160, 555, 1550, 655), "재고 없음\nSOLD_OUT", font, fill="#FCEDEE", outline="#CC6C73")
    arrow(draw, (1080, 605), (1160, 605))
    img.save(path)


def draw_async(path: Path):
    img, draw = create_canvas(1800, 800)
    font = load_diagram_font(30)
    steps = [
        (70, 250, 300, 360, "reserve\n요청"),
        (400, 250, 680, 360, "claim_v2.lua\nstock DECR"),
        (780, 250, 1080, 360, "issue_request\nPROCESSING"),
        (1180, 250, 1460, 360, "outbox_events\nPENDING"),
        (1540, 250, 1760, 360, "Kafka\n발행"),
        (1180, 520, 1460, 630, "Consumer\nuser_tickets 저장"),
        (780, 520, 1080, 630, "SUCCESS\nqueue DONE"),
        (400, 520, 680, 630, "실패 시\n보상 처리"),
    ]
    for xy in steps:
        fill = "#EAF2F8" if "SUCCESS" in xy[4] or "저장" in xy[4] else "#F4F6F9"
        box(draw, xy[:4], xy[4], font, fill=fill)
    for a, b in [(0, 1), (1, 2), (2, 3), (3, 4)]:
        arrow(draw, (steps[a][2], 305), (steps[b][0], 305))
    arrow(draw, (1650, 360), (1320, 520))
    arrow(draw, (1180, 575), (1080, 575))
    arrow(draw, (780, 575), (680, 575))
    img.save(path)


def generate_diagrams(out_dir: Path) -> list[Path]:
    renderers = {
        "dualgate": draw_dualgate,
        "architecture": draw_architecture,
        "sequence": draw_sequence,
        "state": draw_state,
        "erd": draw_erd,
        "queue": draw_queue,
        "async": draw_async,
    }
    paths = []
    for index, spec in enumerate(DIAGRAMS, start=1):
        path = out_dir / f"figure-{index:02d}-{spec.kind}.png"
        renderers[spec.kind](path)
        paths.append(path)
    return paths


def add_figure(doc: Document, image_path: Path):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=4, after=4, line=1.0)
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(6.15))


def parse_body_into_doc(doc: Document, body: str, figure_paths: list[Path]):
    lines = body.splitlines()
    i = 0
    figure_idx = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        if stripped.startswith("```mermaid"):
            if figure_idx < len(figure_paths):
                add_figure(doc, figure_paths[figure_idx])
                figure_idx += 1
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                i += 1
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|---"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_markdown_table(doc, table_lines)
            continue

        heading_match = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading_match:
            hashes, text = heading_match.groups()
            level = max(1, len(hashes) - 1)
            add_heading(doc, text, level)
            i += 1
            continue

        if stripped.startswith("**그림 ") and stripped.endswith("**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(p, before=0, after=8, line=1.15)
            r = p.add_run(stripped.strip("*"))
            set_run_font(r, size=9.5, color=MUTED, bold=True)
            i += 1
            continue

        numbered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline_runs(p, numbered.group(2))
            i += 1
            continue

        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline_runs(p, stripped[2:])
            i += 1
            continue

        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt.startswith("#") or nxt.startswith("|") or nxt.startswith("- ") or nxt.startswith("```"):
                break
            if re.match(r"^\d+\.\s+", nxt):
                break
            paragraph_lines.append(nxt)
            i += 1
        add_rich_paragraph(doc, " ".join(paragraph_lines))


def main():
    markdown = SOURCE_MD.read_text(encoding="utf-8")
    abstract_paragraphs, keywords, body = extract_front_matter(markdown)
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_abstract(doc, abstract_paragraphs, keywords)
    add_toc(doc)

    with tempfile.TemporaryDirectory() as temp:
        figure_paths = generate_diagrams(Path(temp))
        parse_body_into_doc(doc, body, figure_paths)

    finalize_sections(doc)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
